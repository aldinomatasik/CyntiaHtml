import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

IMAGE_WIDTH = Inches(5.5)

def style_cell(cell, fill_hex=None, top_pad=90, bottom_pad=90, left_pad=140, right_pad=140, border_color="CCCCCC"):
    tcPr = cell._tc.get_or_add_tcPr()
    if fill_hex:
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bottom_pad}" w:type="dxa"/><w:left w:w="{left_pad}" w:type="dxa"/><w:right w:w="{right_pad}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)
    if border_color:
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/></w:tcBorders>')
        tcPr.append(borders)

def make_clean_table(doc, title, img_path, note_points, caption=""):
    tbl = doc.add_table(rows=3, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_hdr = tbl.cell(0, 0)
    c_hdr.width = Inches(6.2)
    style_cell(c_hdr, fill_hex="F0F4F8", top_pad=90, bottom_pad=90, left_pad=140, right_pad=140, border_color="B0B0B0")
    
    p_hdr = c_hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    run_hdr = p_hdr.add_run(title)
    run_hdr.bold = True
    run_hdr.font.name = "Calibri"
    run_hdr.font.size = Pt(10)
    run_hdr.font.color.rgb = RGBColor(31, 78, 121)

    c_img = tbl.cell(1, 0)
    c_img.width = Inches(6.2)
    style_cell(c_img, fill_hex="FFFFFF", top_pad=100, bottom_pad=80, left_pad=80, right_pad=80, border_color="CCCCCC")
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=IMAGE_WIDTH)
    else:
        p_img.add_run(f"[Gambar {img_path} tidak ditemukan]")

    if caption:
        p_cap = c_img.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(caption)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9)
        run_cap.font.color.rgb = RGBColor(90, 90, 90)

    c_note = tbl.cell(2, 0)
    c_note.width = Inches(6.2)
    style_cell(c_note, fill_hex="FAFAFA", top_pad=90, bottom_pad=90, left_pad=150, right_pad=150, border_color="CCCCCC")
    
    p_note_title = c_note.paragraphs[0]
    p_note_title.paragraph_format.space_before = Pt(2)
    p_note_title.paragraph_format.space_after = Pt(4)
    r_nlbl = p_note_title.add_run("Note & Analisis Hasil:")
    r_nlbl.bold = True
    r_nlbl.font.name = "Calibri"
    r_nlbl.font.size = Pt(9.5)
    r_nlbl.font.color.rgb = RGBColor(31, 78, 121)

    for pt in note_points:
        p_pt = c_note.add_paragraph()
        p_pt.paragraph_format.line_spacing = 1.15
        p_pt.paragraph_format.space_before = Pt(2)
        p_pt.paragraph_format.space_after = Pt(2)
        p_pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if ":" in pt:
            lead, body = pt.split(":", 1)
            r_lead = p_pt.add_run(f"• {lead.strip()}: ")
            r_lead.bold = True
            r_lead.font.name = "Calibri"
            r_lead.font.size = Pt(9)
            
            r_body = p_pt.add_run(body.strip())
            r_body.font.name = "Calibri"
            r_body.font.size = Pt(9)
        else:
            r_all = p_pt.add_run(f"• {pt.strip()}")
            r_all.font.name = "Calibri"
            r_all.font.size = Pt(9)

    return tbl

def make_clean_code_table(doc, title, img_path, caption=""):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_hdr = tbl.cell(0, 0)
    c_hdr.width = Inches(6.2)
    style_cell(c_hdr, fill_hex="252526", top_pad=80, bottom_pad=80, left_pad=140, right_pad=140, border_color="333333")
    
    p_hdr = c_hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    run_hdr = p_hdr.add_run(title)
    run_hdr.bold = True
    run_hdr.font.name = "Calibri"
    run_hdr.font.size = Pt(9.5)
    run_hdr.font.color.rgb = RGBColor(255, 255, 255)

    c_img = tbl.cell(1, 0)
    c_img.width = Inches(6.2)
    style_cell(c_img, fill_hex="1E1E1E", top_pad=80, bottom_pad=70, left_pad=70, right_pad=70, border_color="333333")
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(3)
    p_img.paragraph_format.space_after = Pt(3)
    
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=IMAGE_WIDTH)
    else:
        p_img.add_run(f"[Gambar {img_path} tidak ditemukan]")

    if caption:
        p_cap = c_img.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(caption)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(8.5)
        run_cap.font.color.rgb = RGBColor(160, 160, 160)

    return tbl

def build_master_report():
    print("Loading base document 'LAPORAN PRAKTIKUM CYNTIA SADE PATILANGI- HTML.bak.docx'...")
    doc = docx.Document('LAPORAN PRAKTIKUM CYNTIA SADE PATILANGI- HTML.bak.docx')
    body = doc._body._element

    # 1. IDENTIFY LANDMARK PARAGRAPHS BY EXACT INDEX
    p_toc = doc.paragraphs[16]._p
    p_dasar = doc.paragraphs[36]._p
    p_tugas_dasar = doc.paragraphs[120]._p
    p_lanjut = doc.paragraphs[149]._p
    p_tugas_lanjut = doc.paragraphs[226]._p

    # 2. STEP 1: CLEAN SECTION D (Remove unfulfilled raw question items 245-254)
    print("Step 1: Removing old raw question copy-paste in Section D...")
    idx_tlanjut = body.index(p_tugas_lanjut)
    elems_d = [body[i] for i in range(idx_tlanjut + 1, len(body)) if body[i].tag.split('}')[-1] != 'sectPr']
    for el in elems_d:
        body.remove(el)
    print(f"Removed {len(elems_d)} raw questions in Section D.")

    # 3. STEP 2: CLEAN SECTION B (Remove broken draft Table 9, 10, headings, duplicate prompts)
    print("Step 2: Removing broken draft fragments in Section B (indices 129 to 160)...")
    idx_tdasar = body.index(p_tugas_dasar)
    idx_lanjut = body.index(p_lanjut)
    elems_b = [body[i] for i in range(idx_tdasar, idx_lanjut)]
    for el in elems_b:
        body.remove(el)
    print(f"Removed {len(elems_b)} old draft elements in Section B.")

    # 4. STEP 3: CLEAN DAFTAR ISI (Remove 19 empty paragraphs)
    print("Step 3: Removing empty paragraphs under DAFTAR ISI...")
    idx_toc = body.index(p_toc)
    idx_dasar = body.index(p_dasar)
    elems_toc = [body[i] for i in range(idx_toc + 1, idx_dasar)]
    for el in elems_toc:
        body.remove(el)
    print(f"Removed {len(elems_toc)} empty paragraphs in DAFTAR ISI.")

    # Helper function to create styled paragraphs
    def make_p(text, bold=False, size=11, before=Pt(4), after=Pt(4), color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = before
        p.paragraph_format.space_after = after
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        return p

    # 5. INSERT COMPLETE DAFTAR ISI (TABLE OF CONTENTS)
    print("Inserting complete DAFTAR ISI...")
    toc_items = [
        ("A. PENGENALAN HTML TINGKAT DASAR", "1", 0, True),
        ("   1. Install dan Jalankan Visual Studio Code", "1", 1, False),
        ("   2. Struktur Dasar HTML", "2", 1, False),
        ("   3. Elemen Head", "3", 1, False),
        ("   4. Elemen Level (Block Level vs Inline Level)", "4", 1, False),
        ("   5. Elemen Text Formatting", "5", 1, False),
        ("   6. Elemen List (Ordered, Unordered, & Description)", "6", 1, False),
        ("   7. Elemen Image", "7", 1, False),
        ("   8. Elemen Audio", "8", 1, False),
        ("   9. Elemen Video", "9", 1, False),
        ("   10. Elemen Table & Nested Table", "10", 1, False),
        ("B. TUGAS HTML DASAR (tugas.html)", "11", 0, True),
        ("   1. Kode Program Visual Studio Code (Baris 1 - 189 Lengkap)", "11", 1, False),
        ("   2. Hasil Output Browser Port 5500 & Analisis", "13", 1, False),
        ("C. PENGENALAN HTML TINGKAT LANJUT", "14", 0, True),
        ("   1. Elemen Form & Ragam Kontrol Input", "14", 1, False),
        ("   2. Elemen Embed", "15", 1, False),
        ("   3. Elemen Object", "16", 1, False),
        ("   4. Elemen Iframe", "17", 1, False),
        ("   5. Tag Semantic HTML5", "18", 1, False),
        ("   6. Visual Studio Code Plugin : Live Server", "19", 1, False),
        ("D. TUGAS HTML LANJUT", "20", 0, True),
        ("   1. Implementasi Tag Semantic HTML5 (semantic.html)", "20", 1, False),
        ("   2. Implementasi Form Registrasi Mahasiswa (form.html)", "23", 1, False),
        ("   3. Publikasi Halaman Web Secara Online Menggunakan Vercel", "26", 1, False)
    ]

    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False

    w_title = Inches(5.6)
    w_page = Inches(0.6)

    for idx, (title, page, level, is_bold) in enumerate(toc_items):
        row = tbl_toc.rows[idx]
        
        c0 = row.cells[0]
        c0.width = w_title
        style_cell(c0, fill_hex=None, top_pad=30, bottom_pad=30, left_pad=50, right_pad=50, border_color=None)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        
        run0 = p0.add_run(title)
        run0.bold = is_bold
        run0.font.name = "Calibri"
        run0.font.size = Pt(9.5)
        if level == 0:
            run0.font.color.rgb = RGBColor(31, 78, 121)

        c1 = row.cells[1]
        c1.width = w_page
        style_cell(c1, fill_hex=None, top_pad=30, bottom_pad=30, left_pad=20, right_pad=20, border_color=None)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        run1 = p1.add_run(page)
        run1.bold = is_bold
        run1.font.name = "Calibri"
        run1.font.size = Pt(9.5)

    p_toc.addnext(tbl_toc._tbl)

    # Page break after TOC
    p_break_toc = doc.add_paragraph()
    p_break_toc.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    tbl_toc._tbl.addnext(p_break_toc._p)

    # 6. FORMAT SECTION A HEADING
    p_dasar_obj = [p for p in doc.paragraphs if p._p == p_dasar][0]
    p_dasar_obj.text = "A. PENGENALAN HTML TINGKAT DASAR"
    p_dasar_obj.paragraph_format.space_before = Pt(12)
    p_dasar_obj.paragraph_format.space_after = Pt(6)
    for r in p_dasar_obj.runs:
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(31, 78, 121)

    # 7. BUILD RECONSTRUCTED SECTION B (TUGAS HTML DASAR)
    print("Inserting brand-new clean Section B with full code screenshots (top to bottom)...")
    # We insert everything before p_lanjut
    def insert_b(xml_el):
        p_lanjut.addprevious(xml_el)

    p_b_title = make_p("B. TUGAS HTML DASAR (tugas.html)", bold=True, size=13, before=Pt(18), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_b(p_b_title._p)

    p_b_desc = make_p("Pada tugas mandiri tingkat dasar ini, mahasiswa membuat berkas tugas.html yang merancang tampilan halaman web profil Program Studi Sistem Informasi ITTelkom Surabaya. Struktur antarmuka disusun secara menyeluruh menggunakan tabel bersarang (nested table) bergaris solid 1px hitam yang proporsional, independen, dan konsisten sesuai acuan modul praktikum.", size=10, before=Pt(2), after=Pt(8), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    insert_b(p_b_desc._p)

    # Subheading 1: Kode Program Lengkap
    p_b_sub1 = make_p("1. Kode Program Visual Studio Code : tugas.html (Baris 1 - 189 Lengkap)", bold=True, size=11, before=Pt(10), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_b(p_b_sub1._p)

    p_b_sub1_note = make_p("Berikut merupakan dokumentasi kode program tugas.html secara lengkap dari baris awal (<!DOCTYPE html>) sampai baris penutup (</html>) yang dibagi ke dalam 3 tangkapan layar editor Visual Studio Code berurutan:", size=9.5, before=Pt(2), after=Pt(6))
    insert_b(p_b_sub1_note._p)

    # Part 1 Code Table
    t_b_c1 = make_clean_code_table(doc, "KODE PROGRAM : tugas.html - BAGIAN 1 (BARIS 1 - 65)", "assets/screenshots/code_tugas_part1.png", "Tangkapan Layar Editor VS Code: tugas.html (Baris 1 - 65: Head, Style CSS, Struktur Tabel, Baris 1 Header Navigasi)")
    insert_b(t_b_c1._tbl)

    p_sp_b1 = make_p("", before=Pt(3), after=Pt(3))
    insert_b(p_sp_b1._p)

    # Part 2 Code Table
    t_b_c2 = make_clean_code_table(doc, "KODE PROGRAM : tugas.html - BAGIAN 2 (BARIS 66 - 125)", "assets/screenshots/code_tugas_part2.png", "Tangkapan Layar Editor VS Code: tugas.html (Baris 66 - 125: Menu Navigasi, Baris 2 Hero Banner, Baris 3 Tiga Kolom Peran)")
    insert_b(t_b_c2._tbl)

    p_sp_b2 = make_p("", before=Pt(3), after=Pt(3))
    insert_b(p_sp_b2._p)

    # Part 3 Code Table
    t_b_c3 = make_clean_code_table(doc, "KODE PROGRAM : tugas.html - BAGIAN 3 (BARIS 126 - 189)", "assets/screenshots/code_tugas_part3.png", "Tangkapan Layar Editor VS Code: tugas.html (Baris 126 - 189: Baris 4 Informasi Footer Tiga Kolom, Baris 5 Hak Cipta)")
    insert_b(t_b_c3._tbl)

    p_sp_b3 = make_p("", before=Pt(6), after=Pt(6))
    insert_b(p_sp_b3._p)

    # Subheading 2: Hasil Output Browser & Analisis
    p_b_sub2 = make_p("2. Hasil Output Browser Port 5500 & Analisis Hasil : tugas.html", bold=True, size=11, before=Pt(12), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_b(p_b_sub2._p)

    tugas_pts = [
        "Struktur Tata Letak (Nested Table): Menggunakan kombinasi tabel bersarang dengan border solid 1px hitam untuk membagi baris secara proporsional dan independen tanpa terpengaruh lebar kolom pada baris lain.",
        "Header & Navigasi: Memuat logo resmi prodi di sebelah kiri, menu navigasi 5 tombol (Tentang Kami, Akademik, Dosen & Staf, Kerja Sama, Blog) di tengah, serta link Admisi berwarna ungu di sebelah kanan.",
        "Hero Banner: Bagian tengah dengan perataan teks center yang memuat judul 'Excellent in Connecting Systems', deskripsi keilmuan SI, dan tombol teks 'Pelajari Lebih Lanjut'.",
        "Tiga Kolom Peran Lulusan: Dibangun dengan tabel bersarang berkolom tiga berukuran 33.33% seimbang untuk mendeskripsikan peran Project Manager, IS Engineer, dan Business & System Analyst.",
        "Footer Informasi Tiga Kolom: Membagi informasi profil prodi (50%), tautan penting dan bermanfaat menggunakan unordered list <ul> (25%), serta info kontak dan media sosial (25%).",
        "Copyright: Menampilkan identitas hak cipta 'Copyright © 2022 Information Systems - Build with ♥ by CODER Team'."
    ]
    t_b_out = make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : tugas.html (PORT 5500)", "assets/screenshots/output_tugas_dasar.png", tugas_pts, "Gambar 2.1: Tampilan Browser tugas.html pada Port 5500 (http://127.0.0.1:5500/tugas.html)")
    insert_b(t_b_out._tbl)

    p_sp_b4 = make_p("", before=Pt(8), after=Pt(8))
    insert_b(p_sp_b4._p)

    # 8. FORMAT SECTION C (PENGENALAN HTML TINGKAT LANJUT)
    print("Formatting Section C...")
    p_lanjut_obj = [p for p in doc.paragraphs if p._p == p_lanjut][0]
    p_lanjut_obj.text = "C. PENGENALAN HTML TINGKAT LANJUT"
    p_lanjut_obj.paragraph_format.space_before = Pt(18)
    p_lanjut_obj.paragraph_format.space_after = Pt(6)
    for r in p_lanjut_obj.runs:
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(31, 78, 121)

    sub_c_map = [
        ("Elemen Form", "1. Elemen Form"),
        ("Elemen Embed", "2. Elemen Embed"),
        ("Elemen Object", "3. Elemen Object"),
        ("Elemen Iframe", "4. Elemen Iframe"),
        ("Tag Semantic HTML5", "5. Tag Semantic HTML5"),
        ("Visual Studio Code Plugin : Live Server", "6. Visual Studio Code Plugin : Live Server")
    ]
    for p in doc.paragraphs:
        for old_txt, new_txt in sub_c_map:
            if p.text.strip() == old_txt:
                p.text = new_txt
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(31, 78, 121)

    # Clean up Section C tables: replace old raw 1080p desktop screenshots with clean cropped images
    clean_c_images = {
        14: 'assets/clean_section_c/table_11_clean.png', # Struktur awal lanjut.html
        15: 'assets/clean_section_c/table_12_clean.png', # Form awal lanjut.html
        16: 'assets/clean_section_c/table_13_clean.png', # Ragam input form
        17: 'assets/clean_section_c/table_14_clean.png', # Embed PDF
        18: 'assets/clean_section_c/table_15_clean.png', # Object PDF
        19: 'assets/clean_section_c/table_16_clean.png', # Iframe YouTube
        20: 'assets/clean_section_c/table_17_clean.png', # Semantic HTML5
    }

    # Style Section A & Section C tables uniformly
    for tbl_idx, tbl in enumerate(doc.tables):
        if tbl in [tbl_toc, t_b_c1, t_b_c2, t_b_c3, t_b_out]:
            continue
        
        # If this is one of Section C tables, swap the image with clean cropped version
        if tbl_idx in clean_c_images and len(tbl.rows) > 1:
            c1_img = tbl.rows[1].cells[0]
            if len(c1_img.paragraphs) > 0:
                p_old_img = c1_img.paragraphs[0]
                for r in list(p_old_img.runs):
                    p_old_img._p.remove(r._r)
                p_old_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_old_img.add_run().add_picture(clean_c_images[tbl_idx], width=IMAGE_WIDTH)

        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        if len(tbl.rows) > 0:
            c0 = tbl.rows[0].cells[0]
            c0.width = Inches(6.2)
            style_cell(c0, fill_hex="F0F4F8", top_pad=70, bottom_pad=70, left_pad=120, right_pad=120, border_color="B0B0B0")
            for p in c0.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(10)
                    r.bold = True
                    r.font.color.rgb = RGBColor(31, 78, 121)
        if len(tbl.rows) > 1:
            c1 = tbl.rows[1].cells[0]
            c1.width = Inches(6.2)
            style_cell(c1, fill_hex="FFFFFF", top_pad=90, bottom_pad=70, left_pad=70, right_pad=70, border_color="CCCCCC")
            for p in c1.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if len(tbl.rows) > 2:
            c2 = tbl.rows[2].cells[0]
            c2.width = Inches(6.2)
            style_cell(c2, fill_hex="FAFAFA", top_pad=90, bottom_pad=90, left_pad=140, right_pad=140, border_color="CCCCCC")
            for p in c2.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)

    # 9. BUILD RECONSTRUCTED SECTION D (TUGAS HTML LANJUT)
    print("Inserting brand-new clean Section D with full code screenshots (top to bottom)...")
    p_tlanjut_obj = [p for p in doc.paragraphs if p._p == p_tugas_lanjut][0]
    p_tlanjut_obj.text = "D. TUGAS HTML LANJUT"
    p_tlanjut_obj.paragraph_format.space_before = Pt(18)
    p_tlanjut_obj.paragraph_format.space_after = Pt(4)
    for r in p_tlanjut_obj.runs:
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(31, 78, 121)

    p_d_desc = doc.add_paragraph()
    p_d_desc.paragraph_format.space_before = Pt(2)
    p_d_desc.paragraph_format.space_after = Pt(8)
    p_d_desc.paragraph_format.line_spacing = 1.15
    p_d_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_dd = p_d_desc.add_run("Pada tugas mandiri tingkat lanjut ini, mahasiswa menyelesaikan implementasi halaman web berbasis tag semantik HTML5, perancangan formulir registrasi interaktif yang komprehensif, serta mempublikasikan seluruh halaman web proyek praktikum ke layanan cloud hosting Vercel:")
    r_dd.font.name = "Calibri"
    r_dd.font.size = Pt(10)

    # -------------------------------------------------------------
    # D.1: semantic.html
    # -------------------------------------------------------------
    p_d_sub1 = doc.add_paragraph()
    p_d_sub1.paragraph_format.space_before = Pt(12)
    p_d_sub1.paragraph_format.space_after = Pt(4)
    r_ds1 = p_d_sub1.add_run("1. Implementasi Tag Semantic HTML5 (semantic.html)")
    r_ds1.bold = True
    r_ds1.font.name = "Calibri"
    r_ds1.font.size = Pt(11)
    r_ds1.font.color.rgb = RGBColor(31, 78, 121)

    p_d_sub1_note = doc.add_paragraph()
    p_d_sub1_note.paragraph_format.space_before = Pt(2)
    p_d_sub1_note.paragraph_format.space_after = Pt(6)
    r_ds1_n = p_d_sub1_note.add_run("Berikut merupakan dokumentasi kode program semantic.html secara lengkap dari baris awal (<!DOCTYPE html>) sampai penutup (</html>) dalam 3 tangkapan layar editor Visual Studio Code berurutan:")
    r_ds1_n.font.name = "Calibri"
    r_ds1_n.font.size = Pt(9.5)

    # Code Parts for semantic.html
    make_clean_code_table(doc, "KODE PROGRAM : semantic.html - BAGIAN 1 (BARIS 1 - 70)", "assets/screenshots/code_semantic_part1.png", "Tangkapan Layar Editor VS Code: semantic.html (Baris 1 - 70: Head, Style CSS, Tag Semantic <header>, dan <nav>)")
    
    p_sp_s1 = doc.add_paragraph()
    p_sp_s1.paragraph_format.space_after = Pt(3)

    make_clean_code_table(doc, "KODE PROGRAM : semantic.html - BAGIAN 2 (BARIS 71 - 137)", "assets/screenshots/code_semantic_part2.png", "Tangkapan Layar Editor VS Code: semantic.html (Baris 71 - 137: Judul Blog, Tag <article> Berita Wisudawan, <iframe> YouTube, Tag <aside> Sidebar)")

    p_sp_s2 = doc.add_paragraph()
    p_sp_s2.paragraph_format.space_after = Pt(3)

    make_clean_code_table(doc, "KODE PROGRAM : semantic.html - BAGIAN 3 (BARIS 138 - 198)", "assets/screenshots/code_semantic_part3.png", "Tangkapan Layar Editor VS Code: semantic.html (Baris 138 - 198: Tag Semantic <footer> Tiga Kolom dan Hak Cipta)")

    p_sp_s3 = doc.add_paragraph()
    p_sp_s3.paragraph_format.space_after = Pt(6)

    # Output Table for semantic.html
    sem_pts = [
        "Tag Semantic HTML5: Menggunakan elemen <header>, <nav>, <main>, <article>, <aside>, dan <footer> sehingga menghasilkan struktur web yang bermakna, rapi, dan ramah mesin pencari (SEO).",
        "Kolom Artikel Berita: Di sisi kiri (<article>), memuat judul berita 'Dua Wisudawan Sistem Informasi Berhasil Meraih Predikat Cumlaude' beserta tiga paragraf narasi prosesi wisuda luring perdana.",
        "Penyematan Video YouTube: Mengintegrasikan rekaman video YouTube Dies Natalis ke-4 ITTelkom Surabaya menggunakan tag <iframe> berukuran 560x315 piksel.",
        "Sidebar Artikel Terkait: Di sisi kanan (<aside>), menampilkan daftar navigasi link untuk Artikel Terbaru dan Artikel Terpopular menggunakan tag <ul> dan link <a> berwarna ungu khas browser.",
        "Footer Halaman: Menampilkan informasi prodi, tautan penting, info kontak, serta hak cipta pada bagian paling bawah halaman."
    ]
    make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : semantic.html (PORT 5500)", "assets/screenshots/output_semantic.png", sem_pts, "Gambar 4.1: Tampilan Browser semantic.html pada Port 5500 (http://127.0.0.1:5500/semantic.html)")

    p_sp_s4 = doc.add_paragraph()
    p_sp_s4.paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # D.2: form.html
    # -------------------------------------------------------------
    p_d_sub2 = doc.add_paragraph()
    p_d_sub2.paragraph_format.space_before = Pt(12)
    p_d_sub2.paragraph_format.space_after = Pt(4)
    r_ds2 = p_d_sub2.add_run("2. Implementasi Form Registrasi Mahasiswa (form.html)")
    r_ds2.bold = True
    r_ds2.font.name = "Calibri"
    r_ds2.font.size = Pt(11)
    r_ds2.font.color.rgb = RGBColor(31, 78, 121)

    p_d_sub2_note = doc.add_paragraph()
    p_d_sub2_note.paragraph_format.space_before = Pt(2)
    p_d_sub2_note.paragraph_format.space_after = Pt(6)
    r_ds2_n = p_d_sub2_note.add_run("Berikut merupakan dokumentasi kode program form.html secara lengkap dari baris awal (<!DOCTYPE html>) sampai penutup (</html>) dalam 3 tangkapan layar editor Visual Studio Code berurutan:")
    r_ds2_n.font.name = "Calibri"
    r_ds2_n.font.size = Pt(9.5)

    # Code Parts for form.html
    make_clean_code_table(doc, "KODE PROGRAM : form.html - BAGIAN 1 (BARIS 1 - 75)", "assets/screenshots/code_form_part1.png", "Tangkapan Layar Editor VS Code: form.html (Baris 1 - 75: Head, Style CSS, Judul Form, Fieldset Biodata: Nama, Readonly NIM, Alamat, Dropdown Tanggal)")

    p_sp_f1 = doc.add_paragraph()
    p_sp_f1.paragraph_format.space_after = Pt(3)

    make_clean_code_table(doc, "KODE PROGRAM : form.html - BAGIAN 2 (BARIS 76 - 155)", "assets/screenshots/code_form_part2.png", "Tangkapan Layar Editor VS Code: form.html (Baris 76 - 155: Dropdown Tanggal, Bulan, Tahun, Radio Gender, File Foto, URL Website, Perguruan Tinggi)")

    p_sp_f2 = doc.add_paragraph()
    p_sp_f2.paragraph_format.space_after = Pt(3)

    make_clean_code_table(doc, "KODE PROGRAM : form.html - BAGIAN 3 (BARIS 156 - 227)", "assets/screenshots/code_form_part3.png", "Tangkapan Layar Editor VS Code: form.html (Baris 156 - 227: Fieldset Info Akun: Email, Username, Password, Ulangi Password, Fieldset Kemampuan Dasar, Tombol Aksi)")

    p_sp_f3 = doc.add_paragraph()
    p_sp_f3.paragraph_format.space_after = Pt(6)

    # Output Table for form.html
    form_pts = [
        "Pengelompokan Fieldset: Menggunakan tag <fieldset> dan <legend> untuk mengelompokkan formulir ke dalam 3 bagian visual: Biodata, Info Akun, dan Kemampuan Dasar.",
        "Perataan Label Form: Menggunakan tabel tanpa garis (borderless table) di dalam fieldset agar label masukan dan tanda titik dua (:) sejajar vertikal secara sempurna.",
        "Kontrol Masukan Biodata: Memuat input teks nama, input readonly untuk NIM (123456789), textarea alamat, 3 dropdown bertingkat <select> untuk tanggal lahir (01-31, Januari-Desember, 1990-2005), radio button jenis kelamin, input file untuk upload foto, URL website, dan perguruan tinggi.",
        "Kontrol Masukan Info Akun: Memuat input bertipe email, text username, password bertipe password, dan ulangi password bertipe password.",
        "Kontrol Kemampuan Dasar & Tombol: Memuat kumpulan checkbox pilihan keahlian jamak (HTML, CSS, JS, PHP, MySQL, Laravel, React Native) serta tiga tombol aksi di bagian bawah: Reset, Simpan, dan Button."
    ]
    make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : form.html (PORT 5500)", "assets/screenshots/output_form.png", form_pts, "Gambar 4.2: Tampilan Browser form.html pada Port 5500 (http://127.0.0.1:5500/form.html)")

    p_sp_f4 = doc.add_paragraph()
    p_sp_f4.paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # D.3: Vercel Hosting
    # -------------------------------------------------------------
    p_d_sub3 = doc.add_paragraph()
    p_d_sub3.paragraph_format.space_before = Pt(14)
    p_d_sub3.paragraph_format.space_after = Pt(4)
    r_ds3 = p_d_sub3.add_run("3. Publikasi Halaman Web Secara Online Menggunakan Vercel")
    r_ds3.bold = True
    r_ds3.font.name = "Calibri"
    r_ds3.font.size = Pt(11)
    r_ds3.font.color.rgb = RGBColor(31, 78, 121)

    p_host = doc.add_paragraph()
    p_host.paragraph_format.line_spacing = 1.25
    p_host.paragraph_format.space_after = Pt(6)
    p_host.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    p_host.add_run("Sesuai instruksi penugasan pada modul praktikum halaman 17, seluruh berkas proyek praktikum ini dipersiapkan dan dipublikasikan secara online menggunakan platform cloud hosting ")
    r_v = p_host.add_run("Vercel")
    r_v.bold = True
    p_host.add_run(" agar dapat diakses secara publik melalui jaringan internet:\n\n")
    
    r_u1 = p_host.add_run("• Domain URL Hosting Vercel (Produksi): ")
    r_u1.bold = True
    r_u2 = p_host.add_run("https://praktikum-paw-html-cyntia.vercel.app/\n\n")
    r_u2.font.color.rgb = RGBColor(0, 32, 96)
    r_u2.bold = True
    r_u2.underline = True

    p_host.add_run("Tahapan Pelaksanaan Deployment ke Vercel:\n1. Persiapan Berkas Proyek: Menyiapkan seluruh berkas web yang telah diselesaikan (index.html sebagai portal navigasi beranda, tugas.html, semantic.html, form.html, serta folder assets/logo.png).\n2. Akses Platform Vercel: Masuk ke dashboard layanan Vercel (https://vercel.com/) menggunakan akun GitHub / Email mahasiswa.\n3. Import Repositori Proyek: Melakukan import repositori proyek atau upload folder praktikum ('cysa html') ke dashboard Vercel.\n4. Konfigurasi Project Settings: Memilih preset static site (Framework Preset: Other) dengan direktori root ('./') sebagai target build.\n5. Verifikasi Publikasi Berhasil: Menekan tombol 'Deploy'. Sistem Vercel melakukan build kilat dalam hitungan detik dan merilis domain produksi publik berstatus aktif (Production) yang dapat diuji langsung oleh dosen penguji maupun asisten laboratorium.")

    # Screenshot of Index Portal on Vercel
    if os.path.exists("assets/screenshots/output_index.png"):
        host_pts = [
            "Halaman Portal Utama (index.html): Berfungsi sebagai landing page utama yang memuat navigasi langsung ke seluruh berkas tugas praktikum (Tugas HTML Dasar, Semantic HTML5, dan Form Registrasi).",
            "Aksesibilitas Global: Halaman web terpublikasi aktif di cloud Vercel dengan protokol HTTPS yang aman dan dapat diuji secara responsif di semua perangkat peramban."
        ]
        make_clean_table(doc, "TAMPILAN PORTAL UTAMA PROYEK (VERCEL DEPLOYMENT)", "assets/screenshots/output_index.png", host_pts, "Gambar 4.3: Tampilan Antarmuka Portal Praktikum Pemrograman Web")

    # SAVE TO ALL MAIN TARGETS
    targets = [
        'LAPORAN_PRAKTIKUM_CYNTIA_SADE_PATILANGI_HTML_FINAL_BERSIH.docx',
        'LAPORAN_PRAKTIKUM_CYNTIA_SADE_PATILANGI_HTML_RAPI.docx',
        'LAPORAN PRAKTIKUM CYNTIA SADE PATILANGI- HTML.docx'
    ]

    for target in targets:
        try:
            doc.save(target)
            print(f"Successfully saved clean master report to: {target}")
        except Exception as e:
            print(f"Notice: Could not save to {target}: {e}")

if __name__ == '__main__':
    build_master_report()
