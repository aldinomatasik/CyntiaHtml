import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

IMAGE_WIDTH = Inches(5.5)

def style_cell(cell, fill_hex=None, top_pad=100, bottom_pad=100, left_pad=150, right_pad=150, border_color="CCCCCC"):
    tcPr = cell._tc.get_or_add_tcPr()
    if fill_hex:
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bottom_pad}" w:type="dxa"/><w:left w:w="{left_pad}" w:type="dxa"/><w:right w:w="{right_pad}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)
    if border_color:
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/></w:tcBorders>')
        tcPr.append(borders)

def make_clean_table(doc, title, img_path, note_points, caption=""):
    tbl = doc.add_table(rows=3, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_hdr = tbl.cell(0, 0)
    c_hdr.width = Inches(6.2)
    style_cell(c_hdr, fill_hex="F0F4F8", top_pad=90, bottom_pad=90, left_pad=140, right_pad=140, border_color="B0B0B0")
    
    p_hdr = c_hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    run_hdr = p_hdr.add_run(title)
    run_hdr.bold = True
    run_hdr.font.name = "Calibri"
    run_hdr.font.size = Pt(10.5)
    run_hdr.font.color.rgb = RGBColor(31, 78, 121)

    c_img = tbl.cell(1, 0)
    c_img.width = Inches(6.2)
    style_cell(c_img, fill_hex="FFFFFF", top_pad=120, bottom_pad=80, left_pad=80, right_pad=80, border_color="CCCCCC")
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=IMAGE_WIDTH)
    else:
        p_img.add_run(f"[Gambar {img_path} tidak ditemukan]")

    if caption:
        p_cap = c_img.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(caption)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(90, 90, 90)

    c_note = tbl.cell(2, 0)
    c_note.width = Inches(6.2)
    style_cell(c_note, fill_hex="FAFAFA", top_pad=100, bottom_pad=100, left_pad=160, right_pad=160, border_color="CCCCCC")
    
    p_note_title = c_note.paragraphs[0]
    p_note_title.paragraph_format.space_before = Pt(2)
    p_note_title.paragraph_format.space_after = Pt(4)
    r_nlbl = p_note_title.add_run("Note & Analisis Hasil:")
    r_nlbl.bold = True
    r_nlbl.font.name = "Calibri"
    r_nlbl.font.size = Pt(10)
    r_nlbl.font.color.rgb = RGBColor(31, 78, 121)

    for pt in note_points:
        p_pt = c_note.add_paragraph()
        p_pt.paragraph_format.line_spacing = 1.2
        p_pt.paragraph_format.space_before = Pt(2)
        p_pt.paragraph_format.space_after = Pt(2)
        p_pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if ":" in pt:
            lead, body = pt.split(":", 1)
            r_lead = p_pt.add_run(f"• {lead.strip()}: ")
            r_lead.bold = True
            r_lead.font.name = "Calibri"
            r_lead.font.size = Pt(9.5)
            
            r_body = p_pt.add_run(body.strip())
            r_body.font.name = "Calibri"
            r_body.font.size = Pt(9.5)
        else:
            r_all = p_pt.add_run(f"• {pt.strip()}")
            r_all.font.name = "Calibri"
            r_all.font.size = Pt(9.5)

    return tbl

def make_clean_code_table(doc, title, img_path, caption=""):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_hdr = tbl.cell(0, 0)
    c_hdr.width = Inches(6.2)
    style_cell(c_hdr, fill_hex="252526", top_pad=90, bottom_pad=90, left_pad=140, right_pad=140, border_color="333333")
    
    p_hdr = c_hdr.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    run_hdr = p_hdr.add_run(title)
    run_hdr.bold = True
    run_hdr.font.name = "Calibri"
    run_hdr.font.size = Pt(10.5)
    run_hdr.font.color.rgb = RGBColor(255, 255, 255)

    c_img = tbl.cell(1, 0)
    c_img.width = Inches(6.2)
    style_cell(c_img, fill_hex="1E1E1E", top_pad=100, bottom_pad=80, left_pad=80, right_pad=80, border_color="333333")
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=IMAGE_WIDTH)
    else:
        p_img.add_run(f"[Gambar {img_path} tidak ditemukan]")

    if caption:
        p_cap = c_img.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        run_cap = p_cap.add_run(caption)
        run_cap.italic = True
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = RGBColor(160, 160, 160)

    return tbl

def build_perfect_report():
    print("Loading base document...")
    doc = docx.Document('LAPORAN PRAKTIKUM CYNTIA SADE PATILANGI- HTML.bak.docx')
    body = doc._body._element

    # ==========================================
    # 1. FIX DAFTAR ISI (TABLE OF CONTENTS)
    # ==========================================
    print("Configuring DAFTAR ISI...")
    toc_p = None
    toc_p_idx = None
    for i, p in enumerate(doc.paragraphs):
        if "DAFTAR ISI" in p.text:
            toc_p = p
            toc_p_idx = i
            break

    # Find the next non-empty paragraph (PENGENALAN HTML TINGKAT DASAR)
    next_p_idx = None
    for i in range(toc_p_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip():
            next_p_idx = i
            break

    # Remove all the empty paragraphs between DAFTAR ISI and PENGENALAN HTML TINGKAT DASAR
    if toc_p_idx is not None and next_p_idx is not None:
        for i in range(next_p_idx - 1, toc_p_idx, -1):
            p_to_del = doc.paragraphs[i]
            p_to_del._p.getparent().remove(p_to_del._p)

    toc_items = [
        ("A. PENGENALAN HTML TINGKAT DASAR", "1", 0, True),
        ("   1. Install dan Jalankan Visual Studio Code", "1", 1, False),
        ("   2. Struktur Dasar HTML", "2", 1, False),
        ("   3. Elemen Head", "3", 1, False),
        ("   4. Elemen Level (Block Level vs Inline Level)", "4", 1, False),
        ("   5. Elemen Text Formatting", "5", 1, False),
        ("   6. Elemen List (Ordered, Unordered, & Description)", "6", 1, False),
        ("   7. Elemen Image", "7", 1, False),
        ("   8. Elemen Audio", "8", 1, False),
        ("   9. Elemen Video", "9", 1, False),
        ("   10. Elemen Table & Nested Table", "10", 1, False),
        ("B. TUGAS HTML DASAR (tugas.html)", "11", 0, True),
        ("   1. Kode Program Visual Studio Code (tugas.html)", "11", 1, False),
        ("   2. Hasil Output Browser Port 5500 & Analisis", "12", 1, False),
        ("C. PENGENALAN HTML TINGKAT LANJUT", "13", 0, True),
        ("   1. Elemen Form & Ragam Kontrol Input", "13", 1, False),
        ("   2. Elemen Embed", "14", 1, False),
        ("   3. Elemen Object", "15", 1, False),
        ("   4. Elemen Iframe", "16", 1, False),
        ("   5. Tag Semantic HTML5", "17", 1, False),
        ("   6. Visual Studio Code Plugin : Live Server", "18", 1, False),
        ("D. TUGAS HTML LANJUT", "19", 0, True),
        ("   1. Implementasi Tag Semantic HTML5 (semantic.html)", "19", 1, False),
        ("   2. Implementasi Form Registrasi Mahasiswa (form.html)", "21", 1, False),
        ("   3. Publikasi Halaman Web Secara Online Menggunakan Vercel", "23", 1, False)
    ]

    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False

    w_title = Inches(5.6)
    w_page = Inches(0.6)

    for idx, (title, page, level, is_bold) in enumerate(toc_items):
        row = tbl_toc.rows[idx]
        
        c0 = row.cells[0]
        c0.width = w_title
        style_cell(c0, fill_hex=None, top_pad=35, bottom_pad=35, left_pad=60, right_pad=60, border_color=None)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        
        run0 = p0.add_run(title)
        run0.bold = is_bold
        run0.font.name = "Calibri"
        run0.font.size = Pt(10)
        if level == 0:
            run0.font.color.rgb = RGBColor(31, 78, 121)

        c1 = row.cells[1]
        c1.width = w_page
        style_cell(c1, fill_hex=None, top_pad=35, bottom_pad=35, left_pad=20, right_pad=20, border_color=None)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        run1 = p1.add_run(page)
        run1.bold = is_bold
        run1.font.name = "Calibri"
        run1.font.size = Pt(10)

    # Insert table directly after DAFTAR ISI paragraph
    toc_p._p.addnext(tbl_toc._tbl)

    # Add a page break paragraph after TOC so Section A begins on next page
    p_break = doc.add_paragraph()
    p_break.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    tbl_toc._tbl.addnext(p_break._p)

    # ==========================================
    # 2. DELETE ALL OLD ELEMENTS IN SECTION B
    # ==========================================
    print("Pruning and rebuilding Section B...")
    # Find exact boundary: between Section A's end and PENGENALAN HTML TINGKAT LANJUT
    idx_lanjut = None
    idx_tugas_dasar = None
    for i, el in enumerate(body):
        txt = ''.join(el.itertext()).strip()
        if 'TUGAS HTML DASAR' in txt and idx_tugas_dasar is None:
            idx_tugas_dasar = i
        if 'PENGENALAN HTML TINGKAT LANJUT' in txt and idx_lanjut is None:
            idx_lanjut = i
            break

    print(f"Section B range: {idx_tugas_dasar} to {idx_lanjut}")
    # Remove all elements from idx_tugas_dasar up to idx_lanjut
    elems_to_remove = [body[i] for i in range(idx_tugas_dasar, idx_lanjut)]
    for el in elems_to_remove:
        body.remove(el)
    print(f"Removed {len(elems_to_remove)} old draft/duplicate elements in Section B!")

    # Find the PENGENALAN HTML TINGKAT LANJUT element to insert before it
    lanjut_el = None
    for el in body:
        txt = ''.join(el.itertext()).strip()
        if 'PENGENALAN HTML TINGKAT LANJUT' in txt:
            lanjut_el = el
            break

    def insert_before(element_xml):
        lanjut_el.addprevious(element_xml)

    def create_p(text, bold=False, size=11, before=Pt(4), after=Pt(4), color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = before
        p.paragraph_format.space_after = after
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        return p

    # Insert clean Section B
    p_b_title = create_p("B. TUGAS HTML DASAR (tugas.html)", bold=True, size=13, before=Pt(18), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_before(p_b_title._p)

    p_b_desc = create_p("Pada tugas mandiri tingkat dasar ini, mahasiswa membuat berkas tugas.html yang merancang tampilan halaman web profil Program Studi Sistem Informasi ITTelkom Surabaya. Antarmuka web disusun menggunakan struktur tabel bersarang (nested table) bergaris solid 1px hitam yang rapi dan konsisten sesuai modul praktikum.", size=10.5, before=Pt(2), after=Pt(8), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    insert_before(p_b_desc._p)

    # Sub-item 1: Kode Program
    p_b_sub1 = create_p("1. Kode Program Visual Studio Code (tugas.html)", bold=True, size=11, before=Pt(10), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_before(p_b_sub1._p)

    t_b_code = make_clean_code_table(doc, "KODE PROGRAM : tugas.html (VISUAL STUDIO CODE)", "assets/screenshots/code_tugas_html.png", "Tangkapan Layar Editor Visual Studio Code: tugas.html (Baris 36 - 82)")
    insert_before(t_b_code._tbl)

    p_sp1 = create_p("", before=Pt(4), after=Pt(4))
    insert_before(p_sp1._p)

    # Sub-item 2: Output Browser
    p_b_sub2 = create_p("2. Hasil Output Browser Port 5500 & Analisis (tugas.html)", bold=True, size=11, before=Pt(12), after=Pt(4), color=RGBColor(31, 78, 121))
    insert_before(p_b_sub2._p)

    tugas_pts = [
        "Struktur Tata Letak (Nested Table): Menggunakan kombinasi tabel bersarang dengan border solid 1px hitam untuk membagi baris secara proporsional dan independen tanpa terpengaruh lebar kolom pada baris lain.",
        "Header & Navigasi: Memuat logo resmi prodi di sebelah kiri, menu navigasi 5 tombol (Tentang Kami, Akademik, Dosen & Staf, Kerja Sama, Blog) di tengah, serta link Admisi berwarna ungu di sebelah kanan.",
        "Hero Banner: Bagian tengah dengan perataan teks center yang memuat judul 'Excellent in Connecting Systems', deskripsi keilmuan SI, dan tombol teks 'Pelajari Lebih Lanjut'.",
        "Tiga Kolom Peran Lulusan: Dibangun dengan tabel bersarang berkolom tiga berukuran 33.33% seimbang untuk mendeskripsikan peran Project Manager, IS Engineer, dan Business & System Analyst.",
        "Footer Informasi Tiga Kolom: Membagi informasi profil prodi (50%), tautan penting dan bermanfaat menggunakan unordered list <ul> (25%), serta info kontak dan media sosial (25%).",
        "Copyright: Menampilkan identitas hak cipta 'Copyright © 2022 Information Systems - Build with ♥ by CODER Team'."
    ]
    t_b_out = make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : tugas.html (PORT 5500)", "assets/screenshots/output_tugas_dasar.png", tugas_pts, "Gambar 2.1: Tampilan Browser tugas.html pada Port 5500 (http://127.0.0.1:5500/tugas.html)")
    insert_before(t_b_out._tbl)

    p_sp2 = create_p("", before=Pt(8), after=Pt(8))
    insert_before(p_sp2._p)

    # ==========================================
    # 3. FIX SECTION C HEADINGS & TABLES
    # ==========================================
    print("Formatting Section C...")
    # Find PENGENALAN HTML TINGKAT LANJUT paragraph
    for p in doc.paragraphs:
        if "PENGENALAN HTML TINGKAT LANJUT" in p.text:
            p.text = "C. PENGENALAN HTML TINGKAT LANJUT"
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(13)
                r.font.color.rgb = RGBColor(31, 78, 121)
            break

    # Number the sub-headings in Section C if needed
    sub_c_map = [
        ("Elemen Form", "1. Elemen Form"),
        ("Elemen Embed", "2. Elemen Embed"),
        ("Elemen Object", "3. Elemen Object"),
        ("Elemen Iframe", "4. Elemen Iframe"),
        ("Tag Semantic HTML5", "5. Tag Semantic HTML5"),
        ("Visual Studio Code Plugin : Live Server", "6. Visual Studio Code Plugin : Live Server")
    ]
    for p in doc.paragraphs:
        for old_txt, new_txt in sub_c_map:
            if p.text.strip() == old_txt:
                p.text = new_txt
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    r.bold = True
                    r.font.name = "Calibri"
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(31, 78, 121)

    # Style Section C tables (all tables currently in the document, keeping their original images)
    for tbl in doc.tables:
        if tbl in [tbl_toc, t_b_code, t_b_out]:
            continue
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Style rows
        if len(tbl.rows) > 0:
            c0 = tbl.rows[0].cells[0]
            c0.width = Inches(6.2)
            style_cell(c0, fill_hex="F0F4F8", top_pad=80, bottom_pad=80, left_pad=120, right_pad=120, border_color="B0B0B0")
            for p in c0.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(10.5)
                    r.bold = True
                    r.font.color.rgb = RGBColor(31, 78, 121)
        if len(tbl.rows) > 1:
            c1 = tbl.rows[1].cells[0]
            c1.width = Inches(6.2)
            style_cell(c1, fill_hex="FFFFFF", top_pad=100, bottom_pad=80, left_pad=80, right_pad=80, border_color="CCCCCC")
            for p in c1.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if len(tbl.rows) > 2:
            c2 = tbl.rows[2].cells[0]
            c2.width = Inches(6.2)
            style_cell(c2, fill_hex="FAFAFA", top_pad=100, bottom_pad=100, left_pad=150, right_pad=150, border_color="CCCCCC")
            for p in c2.paragraphs:
                p.paragraph_format.line_spacing = 1.2
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)

    # ==========================================
    # 4. FIX SECTION D (TUGAS HTML LANJUT)
    # ==========================================
    print("Rebuilding Section D...")
    # Find D. TUGAS HTML LANJUT element
    idx_lanjut_tugas = None
    lanjut_tugas_el = None
    for i, el in enumerate(body):
        txt = ''.join(el.itertext()).strip()
        if 'D. TUGAS HTML LANJUT' in txt:
            idx_lanjut_tugas = i
            lanjut_tugas_el = el
            break

    # Remove all the raw question elements after D. TUGAS HTML LANJUT
    elems_after_lanjut = [body[i] for i in range(idx_lanjut_tugas + 1, len(body)) if body[i].tag.split('}')[-1] != 'sectPr']
    for el in elems_after_lanjut:
        body.remove(el)
    print(f"Removed {len(elems_after_lanjut)} raw module question elements in Section D!")

    # Format the D. TUGAS HTML LANJUT heading
    for p in doc.paragraphs:
        if "D. TUGAS HTML LANJUT" in p.text:
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(13)
                r.font.color.rgb = RGBColor(31, 78, 121)
            break

    # Append narrative description
    p_d_desc = doc.add_paragraph()
    p_d_desc.paragraph_format.space_before = Pt(2)
    p_d_desc.paragraph_format.space_after = Pt(8)
    p_d_desc.paragraph_format.line_spacing = 1.15
    p_d_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_dd = p_d_desc.add_run("Pada tugas mandiri tingkat lanjut ini, mahasiswa menyelesaikan implementasi halaman web berbasis tag semantik HTML5, formulir registrasi interaktif, serta mempublikasikan seluruh halaman web ke layanan cloud hosting Vercel:")
    r_dd.font.name = "Calibri"
    r_dd.font.size = Pt(10.5)

    # Sub-item 1: semantic.html
    p_d_sub1 = doc.add_paragraph()
    p_d_sub1.paragraph_format.space_before = Pt(12)
    p_d_sub1.paragraph_format.space_after = Pt(4)
    r_ds1 = p_d_sub1.add_run("1. Implementasi Tag Semantic HTML5 (semantic.html)")
    r_ds1.bold = True
    r_ds1.font.name = "Calibri"
    r_ds1.font.size = Pt(11)
    r_ds1.font.color.rgb = RGBColor(31, 78, 121)

    make_clean_code_table(doc, "KODE PROGRAM : semantic.html (VISUAL STUDIO CODE)", "assets/screenshots/code_semantic_html.png", "Tangkapan Layar Editor Visual Studio Code: semantic.html (Baris 45 - 91)")

    p_d_sp1 = doc.add_paragraph()
    p_d_sp1.paragraph_format.space_after = Pt(4)

    sem_pts = [
        "Tag Semantic HTML5: Menggunakan elemen <header>, <nav>, <main>, <article>, <aside>, dan <footer> sehingga menghasilkan struktur web yang bermakna, rapi, dan ramah mesin pencari (SEO).",
        "Kolom Artikel Berita: Di sisi kiri (<article>), memuat judul berita 'Dua Wisudawan Sistem Informasi Berhasil Meraih Predikat Cumlaude' beserta tiga paragraf narasi prosesi wisuda luring perdana.",
        "Penyematan Video YouTube: Mengintegrasikan rekaman video YouTube Dies Natalis ke-4 ITTelkom Surabaya menggunakan tag <iframe> berukuran 560x315 piksel.",
        "Sidebar Artikel Terkait: Di sisi kanan (<aside>), menampilkan daftar navigasi link untuk Artikel Terbaru dan Artikel Terpopular menggunakan tag <ul> dan link <a> berwarna ungu khas browser.",
        "Footer Halaman: Menampilkan informasi prodi, tautan penting, info kontak, serta hak cipta pada bagian paling bawah halaman."
    ]
    make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : semantic.html (PORT 5500)", "assets/screenshots/output_semantic.png", sem_pts, "Gambar 4.1: Tampilan Browser semantic.html pada Port 5500 (http://127.0.0.1:5500/semantic.html)")

    p_d_sp2 = doc.add_paragraph()
    p_d_sp2.paragraph_format.space_after = Pt(8)

    # Sub-item 2: form.html
    p_d_sub2 = doc.add_paragraph()
    p_d_sub2.paragraph_format.space_before = Pt(12)
    p_d_sub2.paragraph_format.space_after = Pt(4)
    r_ds2 = p_d_sub2.add_run("2. Implementasi Form Registrasi Mahasiswa (form.html)")
    r_ds2.bold = True
    r_ds2.font.name = "Calibri"
    r_ds2.font.size = Pt(11)
    r_ds2.font.color.rgb = RGBColor(31, 78, 121)

    make_clean_code_table(doc, "KODE PROGRAM : form.html (VISUAL STUDIO CODE)", "assets/screenshots/code_form_html.png", "Tangkapan Layar Editor Visual Studio Code: form.html (Baris 46 - 92)")

    p_d_sp3 = doc.add_paragraph()
    p_d_sp3.paragraph_format.space_after = Pt(4)

    form_pts = [
        "Pengelompokan Fieldset: Menggunakan tag <fieldset> dan <legend> untuk mengelompokkan formulir ke dalam 3 bagian visual: Biodata, Info Akun, dan Kemampuan Dasar.",
        "Perataan Label Form: Menggunakan tabel tanpa garis (borderless table) di dalam fieldset agar label masukan dan tanda titik dua (:) sejajar vertikal secara sempurna.",
        "Kontrol Masukan Biodata: Memuat input teks nama, input readonly untuk NIM (123456789), textarea alamat, 3 dropdown bertingkat <select> untuk tanggal lahir (01-31, Januari-Desember, 1990-2005), radio button jenis kelamin, input file untuk upload foto, URL website, dan perguruan tinggi.",
        "Kontrol Masukan Info Akun: Memuat input bertipe email, text username, password bertipe password, dan ulangi password bertipe password.",
        "Kontrol Kemampuan Dasar & Tombol: Memuat kumpulan checkbox pilihan keahlian jamak (HTML, CSS, JS, PHP, MySQL, Laravel, React Native) serta tiga tombol aksi di bagian bawah: Reset, Simpan, dan Button."
    ]
    make_clean_table(doc, "HASIL OUTPUT TAMPILAN BROWSER : form.html (PORT 5500)", "assets/screenshots/output_form.png", form_pts, "Gambar 4.2: Tampilan Browser form.html pada Port 5500 (http://127.0.0.1:5500/form.html)")

    p_d_sp4 = doc.add_paragraph()
    p_d_sp4.paragraph_format.space_after = Pt(8)

    # Sub-item 3: Hosting Vercel
    p_d_sub3 = doc.add_paragraph()
    p_d_sub3.paragraph_format.space_before = Pt(14)
    p_d_sub3.paragraph_format.space_after = Pt(4)
    r_ds3 = p_d_sub3.add_run("3. Publikasi Halaman Web Secara Online Menggunakan Vercel")
    r_ds3.bold = True
    r_ds3.font.name = "Calibri"
    r_ds3.font.size = Pt(11)
    r_ds3.font.color.rgb = RGBColor(31, 78, 121)

    p_host = doc.add_paragraph()
    p_host.paragraph_format.line_spacing = 1.25
    p_host.paragraph_format.space_after = Pt(6)
    p_host.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    p_host.add_run("Sesuai instruksi penugasan pada modul praktikum halaman 17, seluruh berkas proyek praktikum ini dipersiapkan dan dipublikasikan secara online menggunakan platform cloud hosting ")
    r_v = p_host.add_run("Vercel")
    r_v.bold = True
    p_host.add_run(" agar dapat diakses secara publik melalui jaringan internet:\n\n")
    
    r_u1 = p_host.add_run("• Domain URL Hosting Vercel (Produksi): ")
    r_u1.bold = True
    r_u2 = p_host.add_run("https://praktikum-paw-html-cyntia.vercel.app/\n\n")
    r_u2.font.color.rgb = RGBColor(0, 32, 96)
    r_u2.bold = True
    r_u2.underline = True

    p_host.add_run("Tahapan Pelaksanaan Deployment ke Vercel:\n1. Persiapan Berkas Proyek: Menyiapkan seluruh berkas web yang telah diselesaikan (index.html sebagai portal navigasi beranda, tugas.html, semantic.html, form.html, serta folder assets/logo.png).\n2. Akses Platform Vercel: Masuk ke dashboard layanan Vercel (https://vercel.com/) menggunakan akun GitHub / Email mahasiswa.\n3. Import Repositori Proyek: Melakukan import repositori proyek atau upload folder praktikum ('cysa html') ke dashboard Vercel.\n4. Konfigurasi Project Settings: Memilih preset static site (Framework Preset: Other) dengan direktori root ('./') sebagai target build.\n5. Verifikasi Publikasi Berhasil: Menekan tombol 'Deploy'. Sistem Vercel melakukan build kilat dalam hitungan detik dan merilis domain produksi publik berstatus aktif (Production) yang dapat diuji langsung oleh dosen penguji maupun asisten laboratorium.")

    # Save to multiple targets so user gets the cleanest version regardless of which file they open!
    targets = [
        'LAPORAN_PRAKTIKUM_CYNTIA_SADE_PATILANGI_HTML_FINAL_BERSIH.docx',
        'LAPORAN_PRAKTIKUM_CYNTIA_SADE_PATILANGI_HTML_RAPI.docx',
        'LAPORAN PRAKTIKUM CYNTIA SADE PATILANGI- HTML.docx'
    ]

    for target in targets:
        try:
            doc.save(target)
            print(f"Successfully saved clean report to: {target}")
        except Exception as e:
            print(f"Warning: Could not save to {target}: {e}")

if __name__ == '__main__':
    build_perfect_report()
