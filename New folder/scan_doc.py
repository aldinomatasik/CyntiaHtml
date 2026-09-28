import docx

doc = docx.Document('LAPORAN_PRAKTIKUM_CYNTIA_SADE_PATILANGI_HTML_RAPI.docx')

print('=== ALL TABLES ===')
for i, t in enumerate(doc.tables):
    h = t.rows[0].cells[0].text.strip()
    print(f'Table {i}: {h}')

print('\n=== ALL PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print(f'P{i}: {t}')
